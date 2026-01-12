import tkinter as tk
from tkinter import ttk, messagebox
import csv, os
from datetime import datetime

import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

ARQ_DADOS = "C:/log_logistica/dados.csv"
ARQ_LOGS = "C:/log_logistica/logs.csv"
META = 1000

# =====================================================
# UTILIDADES
# =====================================================

def calcular_indice(sep, conf, erros):
    return sep + conf - (erros * 2)

def status_cor(indice):
    if indice >= META:
        return "verde"
    elif indice >= META * 0.7:
        return "amarelo"
    return "vermelho"

def gerar_id():
    if not os.path.exists(ARQ_DADOS):
        return 1
    with open(ARQ_DADOS) as f:
        ids = [int(l[0]) for l in csv.reader(f) if l]
    return max(ids, default=0) + 1

def registrar_log(acao, detalhe):
    with open(ARQ_LOGS, "a", newline="") as f:
        csv.writer(f).writerow([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            acao,
            detalhe
        ])

# =====================================================
# DADOS
# =====================================================

def carregar_dados():
    if not os.path.exists(ARQ_DADOS):
        return []
    with open(ARQ_DADOS) as f:
        return list(csv.reader(f))

def salvar_dados(dados):
    with open(ARQ_DADOS, "w", newline="") as f:
        csv.writer(f).writerows(dados)

def carregar_logs():
    if not os.path.exists(ARQ_LOGS):
        return []
    with open(ARQ_LOGS) as f:
        return list(csv.reader(f))[-50:]

# =====================================================
# CRUD
# =====================================================

def registrar():
    try:
        sep = int(entry_sep.get())
        conf = int(entry_conf.get())
        erros = int(entry_erros.get())
    except:
        messagebox.showerror("Erro", "Valores inválidos")
        return

    nome = entry_nome.get().strip()
    if not nome:
        messagebox.showerror("Erro", "Nome obrigatório")
        return

    indice = calcular_indice(sep, conf, erros)
    status = status_cor(indice)

    dados = carregar_dados()
    novo_id = gerar_id()

    dados.append([
        novo_id, nome, sep, conf, erros,
        indice, status,
        datetime.now().strftime("%Y-%m-%d %H:%M")
    ])

    salvar_dados(dados)
    registrar_log("REGISTROU", f"{nome} (ID {novo_id})")
    limpar_form()
    atualizar_tudo()

def editar():
    if not entry_id.get():
        return

    dados = carregar_dados()
    for d in dados:
        if str(d[0]) == entry_id.get():
            d[1] = entry_nome.get()
            d[2] = entry_sep.get()
            d[3] = entry_conf.get()
            d[4] = entry_erros.get()
            d[5] = calcular_indice(int(d[2]), int(d[3]), int(d[4]))
            d[6] = status_cor(d[5])
            d[7] = datetime.now().strftime("%Y-%m-%d %H:%M")
            registrar_log("EDITOU", f"{d[1]} (ID {d[0]})")

    salvar_dados(dados)
    limpar_form()
    atualizar_tudo()

def excluir():
    if not entry_id.get():
        return
    if not messagebox.askyesno("Confirmação", "Deseja excluir este registro?"):
        return

    dados = carregar_dados()
    dados = [d for d in dados if str(d[0]) != entry_id.get()]
    salvar_dados(dados)

    registrar_log("EXCLUIU", f"ID {entry_id.get()}")
    limpar_form()
    atualizar_tudo()

# =====================================================
# LOGS
# =====================================================

def abrir_logs():
    win = tk.Toplevel(root)
    win.title("📜 Logs")
    win.geometry("800x300")

    tree = ttk.Treeview(win, columns=("Data", "Ação", "Detalhe"), show="headings")
    for c in ("Data", "Ação", "Detalhe"):
        tree.heading(c, text=c)
        tree.column(c, width=200 if c != "Detalhe" else 380)

    tree.pack(fill="both", expand=True)

    for l in carregar_logs():
        tree.insert("", "end", values=l)

# =====================================================
# FILTROS
# =====================================================

def abrir_filtros():
    win = tk.Toplevel(root)
    win.title("🔎 Filtros")
    win.geometry("900x500")

    resultados = []

    frame = tk.Frame(win)
    frame.pack(pady=10)

    tk.Label(frame, text="Nome ou ID:").grid(row=0, column=0)
    entry_filtro = tk.Entry(frame, width=30)
    entry_filtro.grid(row=0, column=1, padx=5)

    lbl_resumo = tk.Label(win, font=("Arial", 10, "bold"))
    lbl_resumo.pack()

    tree = ttk.Treeview(
        win,
        columns=("ID", "Nome", "Sep", "Conf", "Erros", "Índice", "Status", "Data"),
        show="headings"
    )
    for c in tree["columns"]:
        tree.heading(c, text=c)
        tree.column(c, anchor="center")

    tree.pack(fill="both", expand=True, padx=10)

    tree.tag_configure("verde", background="#c8f7c5")
    tree.tag_configure("amarelo", background="#fff3b0")
    tree.tag_configure("vermelho", background="#f7c5c5")

    def aplicar():
        nonlocal resultados
        resultados = []
        tree.delete(*tree.get_children())

        termo = entry_filtro.get().lower()
        dados = carregar_dados()

        s = c = e = i = 0
        for d in dados:
            if termo in d[1].lower() or termo == str(d[0]):
                resultados.append(d)
                s += int(d[2])
                c += int(d[3])
                e += int(d[4])
                i += int(d[5])
                tree.insert("", "end", values=d, tags=(d[6],))

        lbl_resumo.config(
            text=f"Separação: {s} | Conferência: {c} | Erros: {e} | Índice Total: {i}"
        )

    def exportar():
        if not resultados:
            return

        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "Nome", "Separação", "Conferência", "Erros", "Índice"])

        for d in resultados:
            ws.append(d[:6])

        chart = BarChart()
        data = Reference(ws, min_col=6, min_row=1, max_row=len(resultados)+1)
        cats = Reference(ws, min_col=2, min_row=2, max_row=len(resultados)+1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "H2")

        nome = f"filtro_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(nome)
        messagebox.showinfo("Sucesso", nome)

    tk.Button(frame, text="Aplicar Filtro", command=aplicar).grid(row=0, column=2, padx=5)
    tk.Button(frame, text="Exportar Excel", command=exportar).grid(row=0, column=3)

# =====================================================
# EXCEL GERAL
# =====================================================

def exportar_excel():
    dados = carregar_dados()
    if not dados:
        return

    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Nome", "Separação", "Conferência", "Erros", "Índice"])

    for d in dados:
        ws.append(d[:6])

    chart = BarChart()
    data = Reference(ws, min_col=6, min_row=1, max_row=len(dados)+1)
    cats = Reference(ws, min_col=2, min_row=2, max_row=len(dados)+1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "H2")

    nome = f"geral_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(nome)
    messagebox.showinfo("Exportado", nome)

# =====================================================
# GRÁFICOS
# =====================================================

def atualizar_grafico():
    dados = carregar_dados()
    nomes = [d[1] for d in dados]
    indices = [int(d[5]) for d in dados]

    ax.clear()
    ax.bar(nomes, indices)
    ax.axhline(META, linestyle="--")
    ax.set_title("Desempenho Geral")
    canvas.draw()

def grafico_individual():
    if not entry_id.get():
        return
    for d in carregar_dados():
        if str(d[0]) == entry_id.get():
            plt.figure()
            plt.bar(["Separação", "Conferência", "Erros"],
                    [int(d[2]), int(d[3]), int(d[4])])
            plt.title(d[1])
            plt.show()

# =====================================================
# INTERFACE
# =====================================================

root = tk.Tk()
root.title("Gestão de Logística")
root.geometry("1000x700")

# ---------- FORM ----------
frame_form = tk.LabelFrame(root, text="📋 Dados do Funcionário")
frame_form.pack(fill="x", padx=10, pady=5)

labels = ["ID", "Nome", "Separação", "Conferência", "Erros"]
entries = []

for i, l in enumerate(labels):
    tk.Label(frame_form, text=l).grid(row=i, column=0, sticky="w", pady=2)
    e = tk.Entry(frame_form, width=30)
    e.grid(row=i, column=1, pady=2)
    entries.append(e)

entry_id, entry_nome, entry_sep, entry_conf, entry_erros = entries
entry_id.config(state="readonly")

# ---------- BOTÕES ----------
frame_botoes = tk.LabelFrame(root, text="⚙️ Ações")
frame_botoes.pack(fill="x", padx=10, pady=5)

tk.Button(frame_botoes, text="➕ Registrar", width=20, command=registrar).grid(row=0, column=0, padx=5, pady=5)
tk.Button(frame_botoes, text="✏️ Editar", width=20, command=editar).grid(row=0, column=1, padx=5, pady=5)
tk.Button(frame_botoes, text="🗑️ Excluir", width=20, command=excluir).grid(row=0, column=2, padx=5, pady=5)

tk.Button(frame_botoes, text="📊 Gráfico Funcionário", width=20, command=grafico_individual).grid(row=1, column=0, padx=5)
tk.Button(frame_botoes, text="📜 Logs", width=20, command=abrir_logs).grid(row=1, column=1, padx=5)
tk.Button(frame_botoes, text="🔎 Filtros", width=20, command=abrir_filtros).grid(row=1, column=2, padx=5)

tk.Button(frame_botoes, text="📤 Excel Geral", width=20, command=exportar_excel).grid(row=2, column=1, pady=5)

# ---------- TABELA ----------
frame_tabela = tk.LabelFrame(root, text="📑 Registros")
frame_tabela.pack(fill="both", expand=True, padx=10, pady=5)

tree = ttk.Treeview(
    frame_tabela,
    columns=("ID", "Nome", "Sep", "Conf", "Erros", "Índice", "Status", "Data"),
    show="headings"
)
for c in tree["columns"]:
    tree.heading(c, text=c)
    tree.column(c, anchor="center")

tree.pack(fill="both", expand=True)

tree.tag_configure("verde", background="#c8f7c5")
tree.tag_configure("amarelo", background="#fff3b0")
tree.tag_configure("vermelho", background="#f7c5c5")

def selecionar(event):
    item = tree.selection()
    if not item:
        return
    v = tree.item(item)["values"]

    entry_id.config(state="normal")
    entry_id.delete(0, tk.END)
    entry_id.insert(0, v[0])
    entry_id.config(state="readonly")

    entry_nome.delete(0, tk.END)
    entry_nome.insert(0, v[1])
    entry_sep.delete(0, tk.END)
    entry_sep.insert(0, v[2])
    entry_conf.delete(0, tk.END)
    entry_conf.insert(0, v[3])
    entry_erros.delete(0, tk.END)
    entry_erros.insert(0, v[4])

tree.bind("<<TreeviewSelect>>", selecionar)

# ---------- GRÁFICO ----------
frame_grafico = tk.LabelFrame(root, text="📈 Desempenho Geral")
frame_grafico.pack(fill="x", padx=10, pady=5)

fig, ax = plt.subplots(figsize=(7, 3))
canvas = FigureCanvasTkAgg(fig, master=frame_grafico)
canvas.get_tk_widget().pack(fill="both", expand=True)

def atualizar_tudo():
    tree.delete(*tree.get_children())
    for d in carregar_dados():
        tree.insert("", "end", values=d, tags=(d[6],))
    atualizar_grafico()

def limpar_form():
    for e in entries:
        e.config(state="normal")
        e.delete(0, tk.END)
    entry_id.config(state="readonly")

atualizar_tudo()
root.mainloop()
